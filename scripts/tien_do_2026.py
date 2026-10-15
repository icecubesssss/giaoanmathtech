#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Đọc 2 file TIẾN ĐỘ 2026-2027 (Thầy chốt 15/09/2026) → JSON máy đọc được.

Nguồn (tải từ Drive về inputs/seeds/chuong-trinh-hoc-2026/):
    Tien-do-lop-B-Toan-6-9-chot.xlsx       — lớp B, tuần 1 = CN 21/06/2026, sheet "Tiến độ mới Lớp N"
    Tien-do-hoc-lieu-C-Toan-6-9-chot.xlsx  — lớp C, tuần 1 = 29/06/2026, sheet "Tiến độ khối N - 2026"
Hai file LỆCH số tuần → mỗi tầng giữ số tuần của file mình.

Ra:
    config/tien_do/{B,C}-lop-N.json   — từng buổi/ca (tuần, mạch, chương, bài SGV, loại buổi…)
    config/tien_do/hang-doi.json      — việc cần soạn: (tầng, lớp, môn, chương) → danh sách buổi + thư mục seed
    --check: đối chiếu hàng đợi với seed hiện có (độ phủ theo khối/tầng)

Dùng:
    .venv/bin/python scripts/tien_do_2026.py            # dựng JSON
    .venv/bin/python scripts/tien_do_2026.py --check    # báo độ phủ
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "inputs" / "seeds" / "chuong-trinh-hoc-2026"
FILE_B = SRC / "Tien-do-lop-B-Toan-6-9-chot.xlsx"
FILE_C = SRC / "Tien-do-hoc-lieu-C-Toan-6-9-chot.xlsx"
OUT = ROOT / "config" / "tien_do"
BAN_DO = ROOT / "config" / "ban_do_vd_vdc.json"
LOPS = ("6", "7", "8", "9")

ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5, "VI": 6, "VII": 7, "VIII": 8, "IX": 9, "X": 10}
RE_CHUONG = re.compile(r"Chương\s+([IVX]+)\b")
RE_BAI = re.compile(r"Bài\s+(\d+)\s*\.")

# Loại buổi KHÔNG cần phiếu soạn mới
KHONG_PHIEU = {"Nghỉ"}
# Loại buổi thuộc ôn/thi theo kì (Pha 4), không gắn vào chương
LOAI_KI = {"Luyện tập thi", "Làm đề", "Test", "Ôn – đề", "Kiểm tra", "Chữa đề", "Vận dụng cao", "Tự chọn"}
LOAI_HE = {"Ôn hè", "Khảo sát"}
# Ca B không ghi bài SGV mà luyện lại chương TRƯỚC (suy "chương gần nhất" sai) — gán tay
# sau khi đọc cột "Nội dung ca" (15/10/2026). Khoá: (lớp, tuần, mạch).
CHUONG_TAY = {
    ("6", 11, "dai-so"): 1,   # luyện thứ tự phép tính – luỹ thừa
    ("6", 19, "dai-so"): 2,   # ƯCLN – BCNN thực tế
    ("6", 20, "hinh-hoc"): 4, # chu vi – diện tích thực tế, hình ghép
    ("6", 40, "dai-so"): 7,   # số học HK II: phân số – thập phân – phần trăm
    ("7", 39, "dai-so"): 7,   # đa thức: nghiệm, giá trị
    ("7", 39, "hinh-hoc"): 9, # tam giác cân, đường đồng quy
    ("7", 40, "dai-so"): 7,   # tỉ lệ thức + đa thức tổng hợp
    ("8", 19, "hinh-hoc"): 3, # bài hình tứ giác CK1
    ("8", 21, "dai-so"): 2,   # đa thức – HĐT – phân tích nhân tử
    ("8", 39, "dai-so"): 7,   # lập phương trình
    ("8", 39, "hinh-hoc"): 9, # đồng dạng CK2
    ("8", 40, "dai-so"): 7,   # phân thức · phương trình · hàm số
}
# Tên kì ở file B → mã thư mục (file C đã ghi sẵn gk1/ck1/…)
KI = {"Giữa kì I": "gk1", "Cuối kì I": "ck1", "Giữa kì II": "gk2", "Cuối kì II": "ck2",
      "GK1": "gk1", "CK1": "ck1", "GK2": "gk2", "CK2": "ck2", "Sau CK2": "sau-ck2"}


def _s(x) -> str:
    if x is None:
        return ""
    if isinstance(x, dt.datetime):
        return x.date().isoformat()
    if isinstance(x, float) and x.is_integer():
        return str(int(x))
    return str(x).strip()


def _mach(raw: str) -> tuple[str, bool]:
    """'Đại số (mượn ca Hình)' → ('dai-so', True)."""
    muon = "mượn" in raw
    if raw.startswith("Đại") or raw.startswith("[Đại]"):
        return "dai-so", muon
    if raw.startswith("Hình"):
        return "hinh-hoc", muon
    return ("chung" if raw else ""), muon


def _bai(text: str) -> list[int]:
    return [int(m) for m in RE_BAI.findall(text)]


def _ban_do() -> dict[str, dict[int, dict]]:
    d = json.loads(BAN_DO.read_text(encoding="utf-8"))
    out: dict[str, dict[int, dict]] = {}
    for lop, khoi in d["khoi"].items():
        for c in khoi.get("chuong", []):
            m = re.match(r"chuong-(\d+)-", c.get("slug") or "")
            if m:
                out.setdefault(lop, {})[int(m.group(1))] = c
    return out


def doc_c(lop: str) -> list[dict]:
    wb = openpyxl.load_workbook(FILE_C, data_only=True)
    ws = wb[f"Tiến độ khối {lop} - 2026"]
    hdr = [_s(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    col = {h: i for i, h in enumerate(hdr)}

    def g(r, name):
        for h, i in col.items():
            if h.startswith(name):
                return _s(r[i]) if i < len(r) else ""
        return ""

    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        chuong_raw = g(r, "Chương")
        m = RE_CHUONG.search(chuong_raw)
        mach, muon = _mach(g(r, "Mạch"))
        bai_sgv = g(r, "Bài SGV")
        rows.append({
            "tang": "C", "lop": f"lop-{lop}",
            "tuan": int(float(g(r, "Tuần ở trung tâm") or 0)),
            "buoi": int(float(g(r, "Buổi học tại trung tâm"))) if g(r, "Buổi học tại trung tâm").replace(".", "").isdigit() else None,
            "tu_ngay": g(r, "Từ ngày"), "den_ngay": g(r, "Đến ngày"),
            "luu_y": g(r, "Lưu ý"),
            "mach": mach, "muon_ca_hinh": muon,
            "chuong_so": ROMAN.get(m.group(1)) if m else None,
            "chuong_raw": chuong_raw,
            "thoi_luong": g(r, "Thời lượng"),
            "huong_toi_ki": g(r, "Hướng tới kì"),
            "loai": g(r, "Loại buổi"),
            "noi_dung": g(r, "Bài cụ thể"),
            "hoc_lieu_c": g(r, "Học liệu C"),
            "hoc_lieu_b_bo_sung": g(r, "Học liệu B"),
            "dang_trong_tam": g(r, "Dạng thi trọng tâm"),
            "bai_sgv": bai_sgv, "so_bai": _bai(bai_sgv),
            "ghi_chu": g(r, "Ghi chú"), "link": g(r, "Link"),
        })
    return rows


def doc_b(lop: str, bai_chuong: dict[int, int]) -> list[dict]:
    wb = openpyxl.load_workbook(FILE_B, data_only=True)
    ws = wb[f"Tiến độ mới Lớp {lop}"]
    hdr_row = None
    for i, r in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        if _s(r[0]) == "Tuần":
            hdr_row = i
            hdr = [_s(x) for x in r]
            break
    assert hdr_row, f"không thấy dòng tiêu đề sheet lớp {lop}"
    col = {h: i for i, h in enumerate(hdr)}

    def g(r, name):
        for h, i in col.items():
            if h.startswith(name):
                return _s(r[i]) if i < len(r) else ""
        return ""

    rows = []
    last_chuong: dict[str, int | None] = {}
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        if not r[0]:
            continue
        mach, muon = _mach(g(r, "Mạch"))
        bai_sgv = g(r, "Bài SGV")
        so_bai = _bai(bai_sgv)
        loai = g(r, "Loại ca")
        chuong, suy_ra = None, False
        cs = {bai_chuong[b] for b in so_bai if b in bai_chuong}
        if cs:
            chuong = min(cs)
        elif loai not in LOAI_KI | LOAI_HE | KHONG_PHIEU:
            # Luyện dạng / ôn chương: không ghi bài SGV → lấy chương gần nhất cùng mạch
            chuong, suy_ra = last_chuong.get(mach), True
        tuan = int(re.match(r"\d+", g(r, "Tuần")).group())
        if suy_ra and (lop, tuan, mach) in CHUONG_TAY:
            chuong, suy_ra = CHUONG_TAY[(lop, tuan, mach)], False
        if chuong and not suy_ra:
            last_chuong[mach] = chuong
        rows.append({
            "tang": "B", "lop": f"lop-{lop}",
            "tuan": tuan,
            "tuan_raw": g(r, "Tuần"),
            "han_day": g(r, "Hạn dạy"),
            "mach": mach, "muon_ca_hinh": muon,
            "chuong_so": chuong, "chuong_suy_ra": suy_ra,
            "thoi_luong": g(r, "Thời lượng"),
            "huong_toi_ki": g(r, "Hướng tới kì"),
            "loai": loai,
            "noi_dung": g(r, "Nội dung ca"),
            "bai_sgv": bai_sgv, "so_bai": so_bai, "tiet_sgv": g(r, "Tiết SGV"),
            "dang_trong_tam": g(r, "Dạng thi trọng tâm"),
            "muc_can_chac": g(r, "Mức cần chắc"),
            "ghi_chu": g(r, "Lí do"),
            "tien_do_cu": g(r, "Tiến độ cũ"),
            "link": g(r, "Link"),
        })
    return rows


def _nhom(row: dict, ban_do: dict) -> tuple[str, str] | None:
    """Buổi → (môn, thư mục con dưới lop-x/) hoặc None nếu không cần phiếu."""
    loai = row["loai"]
    if loai in KHONG_PHIEU or not loai:
        return None
    mon = row["mach"] if row["mach"] in ("dai-so", "hinh-hoc") else "dai-so"
    if loai in LOAI_HE or row["huong_toi_ki"] == "Hè":
        return mon, "on-he"
    if loai in LOAI_KI or not row.get("chuong_so"):
        ki = KI.get(row["huong_toi_ki"], re.sub(r"\s+", "-", row["huong_toi_ki"].lower()) or "khac")
        return mon, f"on-tap-kiem-tra/{ki}"
    c = ban_do.get(row["lop"], {}).get(row["chuong_so"])
    slug = c["slug"] if c else f"chuong-{row['chuong_so']:02d}"
    return mon, slug


def dung() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    ban_do = _ban_do()
    hang_doi: dict[str, dict] = collections.OrderedDict()
    for lop in LOPS:
        c_rows = doc_c(lop)
        bai_chuong = {b: r["chuong_so"] for r in c_rows if r["chuong_so"] for b in r["so_bai"]}
        b_rows = doc_b(lop, bai_chuong)
        for tang, rows in (("C", c_rows), ("B", b_rows)):
            (OUT / f"{tang}-lop-{lop}.json").write_text(
                json.dumps({"nguon": (FILE_B if tang == "B" else FILE_C).name, "lop": f"lop-{lop}",
                            "tang": tang, "buoi": rows}, ensure_ascii=False, indent=1),
                encoding="utf-8")
            for r in rows:
                n = _nhom(r, ban_do)
                if not n:
                    continue
                mon, sub = n
                folder = f"inputs/seeds/lop-{lop}/{mon}/lop-{tang.lower()}/{sub}"
                key = f"{tang}|lop-{lop}|{mon}|{sub}"
                item = hang_doi.setdefault(key, {"tang": tang, "lop": f"lop-{lop}", "mon": mon,
                                                 "nhom": sub, "folder": folder, "buoi": []})
                item["buoi"].append({k: r.get(k) for k in
                                     ("tuan", "loai", "noi_dung", "bai_sgv", "huong_toi_ki",
                                      "muon_ca_hinh", "chuong_suy_ra")})
        print(f"lớp {lop}: C {len(c_rows)} buổi · B {len(b_rows)} ca")
    (OUT / "hang-doi.json").write_text(json.dumps(list(hang_doi.values()), ensure_ascii=False, indent=1),
                                       encoding="utf-8")
    print(f"hàng đợi: {len(hang_doi)} nhóm → {OUT / 'hang-doi.json'}")


def check() -> None:
    hd = json.loads((OUT / "hang-doi.json").read_text(encoding="utf-8"))
    tong = collections.Counter()
    co = collections.Counter()
    for it in hd:
        if it["lop"] == "lop-9":
            continue
        k = (it["lop"], it["tang"])
        tong[k] += len(it["buoi"])
        folder = ROOT / it["folder"]
        n = len([p for p in folder.rglob("phieu-*.json")]) if folder.exists() else 0
        co[k] += min(n, len(it["buoi"]))
    for k in sorted(tong):
        print(f"{k[0]} {k[1]}: {co[k]}/{tong[k]} buổi có phiếu (ước theo số file phieu-*.json)")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true")
    a = ap.parse_args()
    check() if a.check else dung()
