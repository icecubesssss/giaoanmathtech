"""Dựng KHUÔN Word `templates/docx/mathtech.docx` (header · footer · watermark · styles)
từ `config/design_tokens.json` — đổi token thì chạy lại, không vẽ tay.

    .venv/bin/python -m scripts.make_docx_template      (hoặc: make docx-template)

Bám bản PDF (`base_handout.tex.j2` + `preamble/_macros.tex.j2`):
  • Lề A4: trên 1.13in · dưới 1.49in · trái/phải 0.75in.
  • Đầu trang: ảnh banner kéo hết bề ngang cột chữ (\\mathtechheader).
  • Chân trang: thanh đỏ footbar 5,5pt → TÊN CÔNG TY đậm 13pt foottitle →
    "Website: www.mathtech.vn" + số trang đậm sát phải (\\mathtechfooter).
  • Watermark: ảnh Watermark_Opacity_20.png xoay 60° (ngược chiều kim đồng hồ),
    căn giữa trang, cỡ 0.8, đậm 45% — NHƯNG dựng theo kiểu Word tự sinh (VML
    `WordPictureWatermark`, mổ từ file "chuẩn hình thức" của trung tâm) để Thầy
    thay/xoá được qua Design → Watermark.
  • Khung mảnh cách mép 8mm + 4 góc chữ L đỏ: MỘT ảnh PNG trong suốt khổ A4 neo
    sau chữ trong header (vẽ bằng Pillow ở 300dpi — y hệt toạ độ TikZ của PDF).
  • Đặt CẢ 3 header (mặc định / trang đầu / trang chẵn) giống nhau — như file mẫu
    của trung tâm — để Thầy có bật "Different first page" hay "Odd & even" trong
    Word thì watermark vẫn còn ở mọi trang.
Style: chỉ TẠO MỚI style tiền tố "MT …", không đè màu lên Heading 1/2 có sẵn của Word
(bài học tex2gdoc: đè style có sẵn ⇒ bấm Heading trong menu ra màu lạ).
"""
from __future__ import annotations

import io
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Inches, Pt, RGBColor
from PIL import Image, ImageDraw

from config import settings
from src.compiler import docx_xml as X

OUT = settings.TEMPLATES_DIR / "docx" / "mathtech.docx"
OUT_NGANG = settings.TEMPLATES_DIR / "docx" / "mathtech-ngang.docx"

A4_W_IN, A4_H_IN = 8.2677, 11.6929
# Lề trên 1.30in (PDF 1.13in): Word không có "headsep" — thân bài bắt đầu ngay dưới đáy
# header nếu header dài hơn lề, nên phải cộng sẵn khoảng hở ~12pt dưới banner. Đo pixel
# 22/10/2026 (pdftoppm 72dpi, trang 2 phiếu Toan9B-Tuan38): banner + chân trang khớp PDF ±2pt.
MARGIN = dict(top=1.30, bottom=1.49, left=0.75, right=0.75)

BODY_FONT = "Times New Roman"
SANS_FONT = "Be Vietnam Pro"      # font tiêu đề của PDF; máy thiếu ⇒ Word dùng altName Arial
MATH_FONT = "Cambria Math"        # GIỮ NGUYÊN: đổi font Equation là ký hiệu thành ô vuông


def tokens() -> dict:
    return json.loads(Path(settings.DESIGN_TOKENS).read_text(encoding="utf-8"))


# ───────────────────────── ảnh phụ ─────────────────────────
def watermark_png(src: Path, opacity: float) -> bytes:
    """Nhân kênh alpha với `opacity` (PDF vẽ ảnh 20% mực với opacity 0.45).
    Làm mờ SẴN trong ảnh thay vì dùng Washout của Word: Washout tẩy tiếp một ảnh
    vốn đã nhạt ⇒ gần như mất hẳn."""
    im = Image.open(src).convert("RGBA")
    a = im.getchannel("A").point(lambda v: round(v * opacity))
    im.putalpha(a)
    buf = io.BytesIO()
    im.save(buf, "PNG", dpi=(300, 300))
    return buf.getvalue()


def frame_png(rule_hex: str, brand_hex: str, dpi: int = 300) -> bytes:
    """Khung trang + 4 góc L — toạ độ chép đúng từ \\AddToShipoutPictureBG của PDF:
    khung 0,5pt màu `rule` cách mép 8mm; mỗi góc hai nét 1,5pt màu `brand` từ 8,5mm
    tới 18mm tính từ mép."""
    W, H = round(A4_W_IN * dpi), round(A4_H_IN * dpi)
    mm = dpi / 25.4
    pt = dpi / 72
    im = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    rgb = lambda h: tuple(int(h[i:i + 2], 16) for i in (0, 2, 4)) + (255,)
    o = 8 * mm
    d.rectangle([o, o, W - o, H - o], outline=rgb(rule_hex), width=max(1, round(0.5 * pt)))
    a, b, lw = 8.5 * mm, 18 * mm, max(1, round(1.5 * pt))
    for sx, sy in ((1, 1), (-1, 1), (1, -1), (-1, -1)):
        x0 = a if sx > 0 else W - a
        y0 = a if sy > 0 else H - a
        x1 = b if sx > 0 else W - b
        y1 = b if sy > 0 else H - b
        d.line([(x0, y0), (x1, y0)], fill=rgb(brand_hex), width=lw)
        d.line([(x0, y0), (x0, y1)], fill=rgb(brand_hex), width=lw)
    buf = io.BytesIO()
    im.save(buf, "PNG", dpi=(dpi, dpi), optimize=True)
    return buf.getvalue()


# ───────────────────────── styles ─────────────────────────
def _rfonts(el, font: str):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.insert(0, rf)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{k}"), font)
    for k in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        rf.attrib.pop(qn(f"w:{k}"), None)


def setup_styles(doc, C: dict):
    st = doc.styles
    # Mặc định tài liệu: Times New Roman 12 (chuẩn file Word của trung tâm), màu mực ink.
    defaults = st.element.find(qn("w:docDefaults"))
    rpr_default = defaults.find(qn("w:rPrDefault")).find(qn("w:rPr"))
    rf = rpr_default.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr_default.insert(0, rf)
    for k in list(rf.attrib):
        del rf.attrib[k]
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{k}"), BODY_FONT)
    for tag, val in (("w:sz", "24"), ("w:szCs", "24")):
        el = rpr_default.find(qn(tag))
        if el is None:
            el = parse_xml(f"<{tag} {nsdecls('w')}/>")
            rpr_default.append(el)
        el.set(qn("w:val"), val)
    # Toán dùng Cambria Math (m:mathPr trong settings sẽ trỏ font này).

    normal = st["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string(C["ink"])
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(4)                  # \parskip 4pt
    pf.line_spacing = 1.15                  # linestretch 1.16 — kiểu "Multiple": dòng có phân số tự giãn
    pf.widow_control = True

    def para_style(name, base="Normal", *, font=None, size=None, bold=None, italic=None,
                   color=None, before=None, after=None, align=None, keep_next=False, outline=None):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = st[base]
        s.quick_style = True
        if font:
            s.font.name = font
            _rfonts(s.element, font)
        if size:
            s.font.size = Pt(size)
        if bold is not None:
            s.font.bold = bold
        if italic is not None:
            s.font.italic = italic
        if color:
            s.font.color.rgb = RGBColor.from_string(color)
        if before is not None:
            s.paragraph_format.space_before = Pt(before)
        if after is not None:
            s.paragraph_format.space_after = Pt(after)
        if align is not None:
            s.paragraph_format.alignment = align
        if keep_next:
            s.paragraph_format.keep_with_next = True
        if outline is not None:
            X.ppr_set(s.element, parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="{outline}"/>'))
        return s

    # Tiêu đề phiếu (\titlefit … 18pt brand) — mức 0 trong Navigation pane.
    para_style("MT Tiêu đề phiếu", font=SANS_FONT, size=18, bold=True, color=C["brand"],
               after=6, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True, outline=0)
    # Đầu chặng ("1. Khám phá") — badge màu chặng do renderer tô trực tiếp.
    para_style("MT Chặng", font=SANS_FONT, size=14, bold=True, color=C["ink"],
               before=12, after=6, keep_next=True, outline=1)
    # \phanmuc — "A. Lý thuyết" (màu chặng, gạch chân mảnh)
    para_style("MT Phần", font=SANS_FONT, size=13, bold=True, color=C["brand"],
               before=12, after=5, keep_next=True, outline=2)
    # \lythuyetmuc — "1. Đường tròn"
    para_style("MT Mục", font=SANS_FONT, size=12, bold=True, color=C["brand"],
               before=8, after=3, keep_next=True, outline=3)
    # \worksheetsection — thanh BTVN nền brand chữ trắng
    s = para_style("MT Mục phụ", font=SANS_FONT, size=12, bold=True, color="FFFFFF",
                   before=13, after=7, keep_next=True, outline=1)
    X.ppr_set(s.element, X.shd(C["brand"]))
    # Chú thích hình (\hinhcap: sans footnotesize nghiêng màu muted)
    para_style("MT Chú thích", font=SANS_FONT, size=9, italic=True, color=C["muted"],
               after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    # Hình căn giữa
    para_style("MT Hình", after=2, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    # Đoạn trong hộp (callout / gợi ý / quy trình): khít hơn thân bài
    para_style("MT Trong hộp", after=2)
    # Nhãn hộp ("VÍ DỤ MẪU", "Gợi ý:"…)
    para_style("MT Nhãn hộp", font=SANS_FONT, size=10, bold=True, after=2, keep_next=True)
    # Chữ nhỏ (dòng "LỚP C", grade label, tiêu đề cột…)
    para_style("MT Nhỏ", font=SANS_FONT, size=9, color=C["muted"], after=2)

    # Header/Footer có sẵn: bỏ khoảng cách để banner/footer nằm đúng chỗ.
    for name in ("Header", "Footer"):
        if name in [x.name for x in st]:
            st[name].paragraph_format.space_after = Pt(0)
            st[name].paragraph_format.line_spacing = 1.0


# ───────────────────────── header / footer ─────────────────────────
def build_header(hdr, C, B, banner: Path, wm_bytes: bytes, frame_bytes: bytes, n: int, text_w_in: float):
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(banner), width=Inches(text_w_in))
    part = hdr.part
    # Watermark VML (ảnh khai trong rels CỦA CHÍNH header này — mỗi header một phần riêng).
    rid_wm, _ = part.get_or_add_image(io.BytesIO(wm_bytes))
    im = Image.open(io.BytesIO(wm_bytes))
    dpi = im.info.get("dpi", (300, 300))[0] or 300
    w_pt = im.width / dpi * 72 * 0.8
    h_pt = im.height / dpi * 72 * 0.8
    p._p.append(parse_xml(X.watermark_vml(rid_wm, w_pt, h_pt, -60, 8565900 + n)))
    # Khung trang + góc L: PNG khổ A4 neo tại (0,0) trang, sau chữ.
    rid_fr, _ = part.get_or_add_image(io.BytesIO(frame_bytes))
    p._p.append(parse_xml(X.anchor_picture_xml(
        rid_fr, round(A4_W_IN * X.EMU_PER_IN), round(A4_H_IN * X.EMU_PER_IN), 0, 0,
        pic_id=900 + n, name="Khung trang")))


def build_footer(ftr, C, B, text_w_in: float):
    ftr.is_linked_to_previous = False
    p1 = ftr.paragraphs[0]
    pf = p1.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # Thanh đỏ dày 5,5pt = viền TRÊN của đoạn tên công ty, cách chữ 4pt.
    X.ppr_set(p1._p, parse_xml(X.border_xml("pBdr", {"top": (5.5, C["footbar"], 4)})))
    r = p1.add_run(B["company_fullname"])
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string(C["foottitle"])

    p2 = ftr.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.tab_stops.clear_all()
    p2.paragraph_format.tab_stops.add_tab_stop(Inches(text_w_in), WD_TAB_ALIGNMENT.RIGHT)
    r = p2.add_run("Website: ")
    r.font.size = Pt(12)
    # Hyperlink gạch chân (\href … \underline)
    rid = ftr.part.relate_to(f"http://{B['website']}", RT.HYPERLINK, is_external=True)
    p2._p.append(parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:rPr>'
        f'<w:color w:val="0000FF"/><w:sz w:val="24"/><w:u w:val="single"/></w:rPr>'
        f'<w:t>{B["website"]}</w:t></w:r></w:hyperlink>'))
    p2.add_run("\t")
    for el in X.field_runs("PAGE", f'<w:b/><w:color w:val="{C["ink"]}"/><w:sz w:val="24"/>'):
        p2._p.append(el)


def build(out: Path, ngang: bool = False) -> Path:
    """`ngang=False`: phiếu/đề A4 dọc CÓ header·footer·watermark.
    `ngang=True` : thuyết minh A4 NGANG, lề 0,8/1,3cm, KHÔNG trang trí — y bản PDF
    (base_thuyetminh.tex.j2 không có đầu/chân trang)."""
    T = tokens()
    C, B = T["colors"], T["branding"]
    assets = settings.ROOT / "assets"
    doc = Document()
    # Thân khuôn để TRỐNG — renderer mở khuôn rồi đổ nội dung vào.
    body = doc.element.body
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)

    sec = doc.sections[0]
    setup_styles(doc, C)
    if ngang:
        from docx.enum.section import WD_ORIENT
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Inches(A4_H_IN), Inches(A4_W_IN)
        sec.top_margin = sec.bottom_margin = Inches(0.8 / 2.54)
        sec.left_margin = sec.right_margin = Inches(1.3 / 2.54)
        sec.header_distance = sec.footer_distance = Inches(0.3 / 2.54)
        normal = doc.styles["Normal"]
        normal.font.size = Pt(11)                      # \documentclass[11pt]
        normal.paragraph_format.space_after = Pt(3)    # \parskip 3pt
    else:
        sec.page_width, sec.page_height = Inches(A4_W_IN), Inches(A4_H_IN)
        sec.top_margin, sec.bottom_margin = Inches(MARGIN["top"]), Inches(MARGIN["bottom"])
        sec.left_margin, sec.right_margin = Inches(MARGIN["left"]), Inches(MARGIN["right"])
        # Banner (cao ~37pt) bắt đầu 44pt từ mép trên — đúng chỗ banner của PDF.
        sec.header_distance = Pt(44)
        # Chân trang: dòng website cách mép dưới ~34pt như PDF (trong khung 8mm, trên góc L).
        sec.footer_distance = Pt(33)
        text_w_in = A4_W_IN - MARGIN["left"] - MARGIN["right"]
        wm = watermark_png(assets / B["watermark_file"], 0.45)
        fr = frame_png(C["rule"], C["brand"])
        banner = assets / B["header_file"]
        sec.different_first_page_header_footer = True
        doc.settings.odd_and_even_pages_header_footer = True
        for n, hdr in enumerate((sec.header, sec.first_page_header, sec.even_page_header)):
            build_header(hdr, C, B, banner, wm, fr, n, text_w_in)
        for ftr in (sec.footer, sec.first_page_footer, sec.even_page_footer):
            build_footer(ftr, C, B, text_w_in)

    settings_el = doc.settings.element
    # Toán: font Cambria Math, công thức riêng dòng căn giữa. (Lưu ý: Word LUÔN thu nhỏ phân
    # số của công thức NẰM GIỮA DÒNG chữ — smallFrac chỉ tác động công thức đứng riêng dòng.)
    mathpr = parse_xml(
        f'<m:mathPr xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
        f'<m:mathFont m:val="{MATH_FONT}"/><m:brkBin m:val="before"/><m:brkBinSub m:val="--"/>'
        f'<m:smallFrac m:val="0"/><m:dispDef/><m:lMargin m:val="0"/><m:rMargin m:val="0"/>'
        f'<m:defJc m:val="centerGroup"/><m:wrapIndent m:val="1440"/><m:intLim m:val="subSup"/>'
        f'<m:naryLim m:val="undOvr"/></m:mathPr>')
    old = settings_el.find(qn("m:mathPr"))
    if old is not None:
        settings_el.remove(old)
    settings_el.append(mathpr)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    _add_font_alt(out)
    print(f"OK → {out}")
    return out


def main() -> int:
    build(OUT)
    build(OUT_NGANG, ngang=True)
    return 0


def _add_font_alt(path: Path):
    """Thêm <w:font w:name="Be Vietnam Pro"><w:altName w:val="Arial"/> vào fontTable."""
    import shutil
    import tempfile
    import zipfile
    tmp = Path(tempfile.mkstemp(suffix=".docx")[1])
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == "word/fontTable.xml":
                s = data.decode("utf-8")
                if SANS_FONT not in s:
                    s = s.replace("</w:fonts>",
                                  f'<w:font w:name="{SANS_FONT}"><w:altName w:val="Arial"/>'
                                  f'<w:charset w:val="00"/><w:family w:val="swiss"/>'
                                  f'<w:pitch w:val="variable"/></w:font></w:fonts>')
                data = s.encode("utf-8")
            zout.writestr(it, data)
    shutil.move(tmp, path)
    path.chmod(0o644)          # mkstemp tạo 600 — để 644 như file thường


if __name__ == "__main__":
    sys.exit(main())
