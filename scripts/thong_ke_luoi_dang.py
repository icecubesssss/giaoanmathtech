#!/usr/bin/env python3
"""Thống kê dạng từ LƯỚI DẠNG gắn tay (inputs/refs/de-thi/lop-9/luoi-dang/<ky>.json).

Khác `thong_ke_dang_de.py` (đếm từ khoá, ~87% tin cậy, SÓT dạng lời văn như "lập BPT"):
lưới dạng do người đọc từng đề và gắn mã, nên con số dùng được để CHỐT KHUNG đề.

Ra: outputs/thong-ke-dang-de/Thong-ke-dang-<KY>-Toan9-luoi-tay.xlsx + .md
    - % số đề có dạng (cả 2 năm / riêng năm mới nhất), điểm TB khi có
    - phủ theo quận/huyện, đề gần trùng nhau

Dùng:
    .venv/bin/python scripts/thong_ke_luoi_dang.py            # mặc định gk1
    .venv/bin/python scripts/thong_ke_luoi_dang.py --ky gk1
"""
from __future__ import annotations

import argparse
import collections
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
LUOI = ROOT / "inputs" / "refs" / "de-thi" / "lop-9" / "luoi-dang"
OUT = ROOT / "outputs" / "thong-ke-dang-de"

# Nhóm hiển thị theo khung đề GK1 (thứ tự = thứ tự bài trong đề trường)
NHOM = [
    ("I. Giải PT – BPT – hệ", ["PT_TICH", "PT_MAU", "PT_BAC1", "BPT_GIAI", "BPT_NB", "BDT_SS",
                               "HE_GIAI", "HE_ANPHU", "HE_THAMSO", "PT2AN"]),
    ("II. Toán lời văn", ["LAP_HE", "LAP_BPT", "SOSANH_TT"]),
    ("III. Lượng giác", ["LG_THUCTE", "LG_BIEUTHUC", "LG_MAYTINH"]),
    ("IV. Hình tam giác vuông", ["TGV_TINH", "TGV_CM", "TGV_VDC", "DTRON"]),
    ("V. Câu cuối", ["TOIUU_TT", "VDC_DS"]),
    ("Khác", ["TN", "CAN", "KHAC"]),
]


def thong_ke(du_lieu: dict) -> dict:
    de = du_lieu["de"]
    nam_moi = max(e["nam"] for e in de)
    kq = {}
    for ma in du_lieu["ma_dang"]:
        co = [e for e in de if any(ma in b["dang"] for b in e["bai"])]
        co_moi = [e for e in co if e["nam"] == nam_moi]
        # điểm của dạng trong 1 đề = tổng điểm các bài có dạng, chia đều cho số dạng của bài
        diem = [sum(b["diem"] / len(b["dang"]) for b in e["bai"] if ma in b["dang"]) for e in co]
        so_bai = sum(1 for e in co for b in e["bai"] if ma in b["dang"])
        kq[ma] = {
            "so_de": len(co), "pt": round(100 * len(co) / len(de)),
            "so_de_moi": len(co_moi),
            "pt_moi": round(100 * len(co_moi) / sum(1 for e in de if e["nam"] == nam_moi)),
            "diem_tb": round(sum(diem) / len(diem), 2) if diem else 0,
            "so_bai": so_bai,
            "truong": [f"{e['truong']} {e['nam'][2:4]}-{e['nam'][7:]}" for e in co],
        }
    return {"nam_moi": nam_moi, "n": len(de), "n_moi": sum(1 for e in de if e["nam"] == nam_moi),
            "dang": kq}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--ky", default="gk1")
    a = ap.parse_args()
    du_lieu = json.loads((LUOI / f"{a.ky}.json").read_text(encoding="utf-8"))
    tk = thong_ke(du_lieu)
    ten = du_lieu["ma_dang"]
    de = du_lieu["de"]
    quan = collections.Counter(e["quan"] for e in de)
    trung = [(e["truong"], e["nam"], e["ghi_chu"]) for e in de
             if "TRÙNG" in e.get("ghi_chu", "") or "trùng" in e.get("ghi_chu", "").lower()]

    md = [f"# Thống kê dạng đề {a.ky.upper()} Toán 9 Hà Nội — gắn tay {tk['n']} đề",
          "",
          f"Nguồn: {tk['n']} đề {a.ky.upper()} Hà Nội (toàn bộ đề trên toanmath, đã bỏ đề lệch KNTT); "
          f"riêng năm {tk['nam_moi']}: {tk['n_moi']} đề. Mỗi bài trong đề được đọc và gắn dạng bằng tay.",
          "", "| Nhóm | Dạng | Số đề | % (cả 2 năm) | % năm " + tk["nam_moi"] + " | Điểm TB khi có |",
          "|---|---|---|---|---|---|"]
    for nhom, mas in NHOM:
        for ma in mas:
            if ma not in tk["dang"]:
                continue
            d = tk["dang"][ma]
            md.append(f"| {nhom} | {ten[ma]} | {d['so_de']} | {d['pt']}% | {d['pt_moi']}% | {d['diem_tb']} |")
    md += ["", f"## Phủ theo quận/huyện cũ ({len(quan)} quận/huyện)", ""]
    md += [f"- {q}: {n} đề" for q, n in quan.most_common()]
    md += ["", "## Đề gần trùng nhau (không đếm là bằng chứng độc lập)", ""]
    md += [f"- {t} {n}: {g}" for t, n, g in trung]
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / f"Thong-ke-dang-{a.ky.upper()}-Toan9-luoi-tay.md").write_text("\n".join(md) + "\n", encoding="utf-8")

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print("thiếu openpyxl — chỉ ra .md")
        return 0
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tổng hợp"
    ws.append([f"THỐNG KÊ DẠNG ĐỀ {a.ky.upper()} TOÁN 9 HÀ NỘI — GẮN TAY {tk['n']} ĐỀ"])
    ws.append([f"Năm {tk['nam_moi']}: {tk['n_moi']} đề. Nguồn: toàn bộ đề trên toanmath, đã bỏ đề lệch KNTT."])
    ws.append([])
    ws.append(["Nhóm", "Dạng", "Số đề", "% cả 2 năm", f"% năm {tk['nam_moi']}", "Điểm TB khi có", "Trường"])
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for nhom, mas in NHOM:
        for ma in mas:
            if ma in tk["dang"]:
                d = tk["dang"][ma]
                ws.append([nhom, ten[ma], d["so_de"], d["pt"], d["pt_moi"], d["diem_tb"], "; ".join(d["truong"])])
    ws["A1"].font = Font(bold=True, size=13)
    for col, w in zip("ABCDEFG", (24, 52, 8, 11, 13, 13, 120)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Lưới từng đề")
    mas = [m for _, ms in NHOM for m in ms]
    ws2.append(["Trường", "Năm", "Quận/huyện cũ"] + mas + ["Ghi chú"])
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(text_rotation=90 if c.column > 3 else 0)
    for e in sorted(de, key=lambda e: (e["nam"], e["quan"], e["truong"])):
        co = {m for b in e["bai"] for m in b["dang"]}
        ws2.append([e["truong"], e["nam"], e["quan"]] + ["x" if m in co else "" for m in mas]
                   + [e.get("ghi_chu", "")])
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["C"].width = 22

    ws3 = wb.create_sheet("Từng bài")
    ws3.append(["Trường", "Năm", "Quận", "Bài", "Điểm", "Dạng", "Ghi chú"])
    for e in de:
        for b in e["bai"]:
            ws3.append([e["truong"], e["nam"], e["quan"], b["vi_tri"], b["diem"],
                        ", ".join(ten[m] for m in b["dang"]), b.get("ghi_chu", "")])
    ws3.column_dimensions["A"].width = 26
    ws3.column_dimensions["F"].width = 70
    ws3.column_dimensions["G"].width = 45
    p = OUT / f"Thong-ke-dang-{a.ky.upper()}-Toan9-luoi-tay.xlsx"
    wb.save(p)
    print(p.relative_to(ROOT))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
