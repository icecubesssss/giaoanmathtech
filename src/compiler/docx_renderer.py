"""Phiếu (LessonPackage) → bản Word (.docx): mở KHUÔN `templates/docx/mathtech.docx`
(header · footer · watermark · style — xem scripts/make_docx_template.py) rồi đổ nội
dung theo ĐÚNG thứ tự và quy tắc của bản PDF (`_blocks.j2`, `components/*.tex.j2`).

Ba tầng:
  1. Phiếu → cây khối IR (`docx_text`): đọc chữ/lệnh LaTeX dàn trang, hộp, bài, hình.
  2. Gom TẤT CẢ công thức → `omml.convert_many` (1 lần pandoc); gom TẤT CẢ hình →
     `tikz_raster.render_many` (1 lần Tectonic). Ghép lại theo khoá, không theo thứ tự.
  3. IR → Word (python-docx + OOXML viết tay ở `docx_xml`).

Không vứt gì im lặng: lệnh lạ / công thức hỏng / hình hỏng ⇒ `report.issues`, cổng
`docx_gate` chặn ghi file.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field, replace
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from config import settings
from src.compiler import docx_xml as X
from src.compiler.docx_text import (Box, Cell, Ctx, DMath, Fig, Fmt, Img, M, P, PageBreak, Reader,
                                    T, Tab, Tbl, WriteLines, color_hex, dots, set_palette, stars)
from src.compiler.jinja_renderer import (_texify, fig_side, load_tokens, mcq_cols, split_reflection,
                                         tach_loi_giai, tach_y_con)
from src.compiler.omml import M_NS, convert_many, is_only_space

TEMPLATE = settings.TEMPLATES_DIR / "docx" / "mathtech.docx"
TEXT_W_IN = 6.77
SANS = "Be Vietnam Pro"

_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
          "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


@dataclass
class Report:
    issues: list = field(default_factory=list)      # lỗi CHẶN ghi file
    counts: dict = field(default_factory=dict)      # số đếm nguồn để cổng đối chiếu


# ═══════════════════════════════ 1. Phiếu → IR ═══════════════════════════════
class LessonIR:
    """Dựng cây khối cho phiếu HS (show_solution=False) hoặc Sổ tay GV (True)."""

    def __init__(self, lesson, tokens: dict, show_solution: bool = False, lesson_dir: Path | None = None):
        self.L = lesson
        self.C = tokens["colors"]
        self.B = tokens["branding"]
        self.show = show_solution
        self.dir = lesson_dir
        self.ctx = Ctx(show_solution=show_solution)
        self.rd = Reader(self.ctx)
        self.out: list = []
        self.n_problem = 0
        self.stagecolor = self.C["brand"]

    # ---------- tiện ích ----------
    def col(self, spec: str) -> str:
        return color_hex(spec, self.ctx) or self.C["ink"]

    def blocks(self, s: str, where: str, fmt: Fmt = Fmt(), align="justify") -> list:
        self.ctx.where = where
        return self.rd.blocks(_texify_light(s), fmt, align)

    def inl(self, s: str, where: str, fmt: Fmt = Fmt()) -> list:
        self.ctx.where = where
        return self.rd.inlines(_texify_light(s), fmt)

    def img_path(self, p: str) -> str:
        if not p:
            return p
        q = Path(p)
        if not q.is_absolute() and self.dir is not None:
            q = self.dir / q
        return str(q)

    # ---------- toàn phiếu ----------
    def build(self) -> list:
        L = self.L
        self.title_card()
        if L.theme != "de_thi":
            self.out.append(Fig(snippet="\\begin{minipage}{\\linewidth}\\journeybar{0}\\end{minipage}",
                                max_w_frac=1.0))
        for k, stage in enumerate(L.stages):
            if L.theme == "de_thi":
                self.plain_stage(stage)
            else:
                self.stage(stage)
            if k < len(L.stages) - 1:
                self.out.append(P([], style=None, border_bottom=self.col("brand!40"),
                                  space_before=6, space_after=10))
        return self.out

    def title_card(self):
        L, C = self.L, self.C
        badge = lambda txt, bg, fg="FFFFFF": T(f"\u00A0{txt}\u00A0", Fmt(bold=True, sans=True, size=9,
                                                                       color=fg, shade=bg))
        row1: list = []
        if L.class_tier:
            row1 += [badge(f"LỚP {L.class_tier}", C["ink"]), T("\u2002")]
        if L.eyebrow:
            row1 += [badge(L.eyebrow, C["brand"])]
        else:
            row1 += [T("PHIẾU HỌC TẬP", Fmt(bold=True, sans=True, size=12, color=C["ink"]))]
        if L.grade_label:
            row1 += [Tab(), T(L.grade_label, Fmt(sans=True, size=10, color=C["muted"]))]
        # PDF in `lesson.title | upper` thô (không qua filter tex) ⇒ ở đây cũng chữ thô.
        title = P([T(L.title.upper(), Fmt())], style="MT Tiêu đề phiếu", align="left")
        name = P([T("Họ và tên:\u00A0", Fmt(sans=True)), T(dots(6.0), Fmt(sans=True)), Tab(),
                  T("Lớp:\u00A0", Fmt(sans=True)), T(dots(3.2), Fmt(sans=True))], align="left")
        self.out.append(Box([P(row1, align="left", space_after=4), title, name],
                            frame=self.col("brand!40"), back=C["neutral"], frame_pt=0.8, pad_pt=7))

    def stage_heading(self, stage, color: str, icon: str):
        badge = Img(snippet=f"\\stagebadge{{{_stage_color_name(stage)}}}{{{icon}}}", h_pt=20)
        self.out.append(P([badge, T("\u2002"), T(f"{stage.number}.\u2003")]
                          + self.inl(stage.title, f"tiêu đề chặng {stage.number}"),
                          style="MT Chặng", align="left", keep_next=True, border_bottom=color))

    def stage(self, stage):
        icon = {"review": "\\iconreview", "concept": "\\iconconcept", "practice1": "\\iconpractice",
                "practice2": "\\iconapply", "reflection": "\\iconreflect"}[stage.kind]
        self.stagecolor = self.C[_stage_color_name(stage)]
        self.stage_heading(stage, self.stagecolor, icon)
        if stage.kind == "reflection":
            parts = split_reflection(stage.blocks)
            self.render_blocks(parts["tong_ket"], stage)
            self.stagecolor = self.C["brand"]
            if parts["btvn"]:
                self.section("BÀI TẬP VỀ NHÀ")
                self.render_blocks(parts["btvn"], stage)
            if parts["mo_rong"]:
                self.section("BÀI TẬP MỞ RỘNG (tự chọn)")
                self.render_blocks(parts["mo_rong"], stage)
        else:
            self.render_blocks(stage.blocks, stage)

    def plain_stage(self, stage):
        """Chặng 'trần' của đề kiểm tra (components/plain.tex.j2): không hộp, không icon;
        tiêu đề rỗng ⇒ không thanh tiêu đề."""
        self.stagecolor = self.C["brand"]
        if stage.kind == "reflection":
            parts = split_reflection(stage.blocks)
            if stage.title:
                self.section(stage.title)
            self.render_blocks(parts["tong_ket"], stage)
            if parts["btvn"]:
                self.section("BÀI TẬP VỀ NHÀ")
                self.render_blocks(parts["btvn"], stage)
            if parts["mo_rong"]:
                self.section("BÀI TẬP MỞ RỘNG (tự chọn)")
                self.render_blocks(parts["mo_rong"], stage)
        else:
            if stage.title:
                self.section(stage.title)
            self.render_blocks(stage.blocks, stage)

    def section(self, title: str):
        """\\worksheetsection — thanh nền brand chữ trắng."""
        self.out.append(P(self.inl(title, "tiêu đề mục"), style="MT Mục phụ", align="left", keep_next=True))

    def render_blocks(self, blocks, stage):
        started = False
        for i, b in enumerate(blocks):
            if not started and b.type == "writelines" and not b.count:
                continue                      # writelines rỗng ở đầu = khoảng trắng chết
            started = True
            self.block(b, f"chặng {stage.number} khối {i + 1}")

    # ---------- từng khối ----------
    def block(self, b, where: str):
        t = b.type
        C = self.C
        if t == "para":
            if b.variant == "heading":
                self.out.append(P(self.inl(b.text, where, Fmt(color=self.stagecolor)),
                                  style="MT Mục", align="left", keep_next=True))
            elif b.variant == "part":
                self.out.append(P(self.inl(b.text, where, Fmt(color=self.stagecolor)),
                                  style="MT Phần", align="left", keep_next=True,
                                  border_bottom=self.stagecolor))
            else:
                self.out.extend(self.blocks(b.text, where))
        elif t == "deflist":
            for it in b.items:
                inl = self.inl(it.term, where, Fmt(bold=True)) + [T(" ")] + self.inl(it.desc, where)
                self.out.append(P(inl, align="justify", indent_pt=14, hanging_pt=14, space_after=2))
        elif t == "math":
            self.out.append(DMath(_texify_light(b.latex)))
        elif t == "noted":
            self.out.append(self.noted(b, where))
        elif t == "opener":
            self.out.append(self.opener(b, where))
        elif t == "writelines":
            if b.variant == "oly":
                self.out.append(Fig(snippet=f"\\begin{{minipage}}{{\\linewidth}}\\olyline[{b.count}]\\end{{minipage}}"))
            else:
                self.out.append(WriteLines(b.count, b.variant or ""))
        elif t == "problem":
            self.problem(b, where)
        elif t == "table":
            self.table(b, where)
        elif t == "mindmap":
            self.mindmap(b, where)
        elif t == "figure":
            self.out.extend(self.figure_block(b, where))
        else:
            self.ctx.where = where
            self.ctx.lạ(f"khối {t}")

    def noted(self, b, where) -> Box:
        C = self.C
        v = b.variant
        target = "ĐÍCH THI VÀO 10" if "Lớp 9" in (self.L.grade_label or "") else "ĐÍCH ĐỀ KIỂM TRA"
        spec = {
            "example": ("info!50", "\u270E  VÍ DỤ MẪU", C["info"], "FFFFFF"),
            "trap": ("trap!55", "\u261E  BẪY ĐIỂM", C["trap"], "FFFFFF"),
            "target": ("gold!75", f"\u2605  {target}", C["gold"], C["ink"]),
            "tip": ("stage2!45", "\u261B  MẸO", C["stage2"], "FFFFFF"),
        }.get(v)
        if b.figure is not None and (b.figure.tikz or b.figure.image):
            de, lg = tach_loi_giai(b.text)
            inner = self.blocks(de, where) + [self.fig_of(b.figure, where, force=0.40, max_h=3.0)]
            inner += self.blocks(lg, where) if lg else []
        else:
            inner = self.blocks(b.text, where)
        if spec:
            frame, title, bg, fg = spec
            return Box(inner, frame=self.col(frame), title=title, title_bg=bg, title_fg=fg)
        return Box(inner, frame=C["rule"], frame_pt=0.4, pad_pt=8)

    def opener(self, b, where) -> Box:
        C = self.C
        text = self.blocks(b.text, where)
        if b.tikz or b.image:
            fig = Fig(snippet=b.tikz, image=self.img_path(b.image), max_w_frac=1.0, max_h_cm=3.4)
            inner = [Tbl([[Cell(text), Cell([fig])]], widths=[0.63, 0.34], borders=False,
                         cell_pad_pt=0, align="left")]
        else:
            inner = text
        return Box(inner, frame=self.col("stage1!55"), title="\u270C  MỞ MÀN THỰC TẾ", title_bg=C["stage1"])

    def fig_of(self, f, where, force: float | None = None, max_h: float | None = None,
               max_w: float = 1.0) -> Fig:
        cap = self.inl(f.caption, where) if f.caption else []
        if f.image:
            return Fig(image=self.img_path(f.image), caption=cap, max_w_frac=max_w,
                       force_w_frac=force, max_h_cm=max_h)
        return Fig(snippet=f.tikz, caption=cap, max_w_frac=max_w, force_w_frac=force, max_h_cm=max_h)

    def label_runs(self, b) -> list:
        runs = [T(_plain_label(b.label), Fmt(bold=True))]
        lv = b.level or {"onclass": 2, "btvn": 1, "extend": 3}.get(b.tier, 0)
        if lv:
            runs += [T("\u00A0")] + stars(lv)
        return runs + [T(" ")]

    def problem(self, b, where):
        self.n_problem += 1
        lead = self.label_runs(b)
        if b.layout_moi:
            de, ycon = tach_y_con(b.statement)
            body = self.blocks(de, where)
            _prepend(body, lead)
            pos = fig_side(b)
            if pos == "below":
                thoang = self.L.trinh_bay == "thoang"
                body.append(self.fig_of(b.figure, where, force=0.5 if thoang else 0.44,
                                        max_h=6.5 if thoang else 5.0))
            if ycon:
                body += self.blocks(ycon, where)
            if b.options:
                n = mcq_cols(b.options)
                cells = [Cell([P([T(f"{'ABCDEFGH'[k]}.", Fmt(bold=True)), T("\u00A0")] + self.inl(o, where),
                                align="left", space_after=2)]) for k, o in enumerate(b.options)]
                rows = [cells[k:k + n] for k in range(0, len(cells), n)]
                for r in rows:
                    while len(r) < n:
                        r.append(Cell([]))
                body.append(Tbl(rows, widths=[1 / n] * n, borders=False, cell_pad_pt=1, align="left"))
        else:
            stmt = b.statement
            if "[[/wrap]]" in stmt:
                body = self.wrap_problem(stmt, lead, where)
            else:
                body = self.blocks(stmt, where)
                _prepend(body, lead)
        if body and isinstance(body[0], P):
            body[0].space_before = 8           # \par\addvspace{3mm} trước mỗi bài
        self.out.extend(body)
        if b.video:
            self.out.append(P([Img(path=f"qr:{b.video}", h_pt=27)], align="right", space_after=0))
            self.out.append(P([T("\u261B video lời giải", Fmt(sans=True, size=7, color=self.C["muted"]))],
                              align="right"))
        for _ in range(b.writelines or 0):
            self.out.append(WriteLines(1))
        if b.quy_trinh:
            inner = [P([T("Quy trình giải:", Fmt(bold=True, sans=True, color=self.C["stage3"]))],
                       style="MT Trong hộp", align="left")]
            for k, st in enumerate(b.quy_trinh):
                inner += self.blocks(st, where)
            self.out.append(Box(inner, frame=self.col("stage3!45"), back=self.col("stage3!6"),
                                frame_pt=0.5, pad_pt=5))
        if b.hints:
            inner = [P([T("Gợi ý:", Fmt(italic=True, color=self.C["brand"]))], style="MT Trong hộp",
                       align="left")]
            for h in b.hints:
                inner.append(P([T("\u2022 ")] + self.inl(h, where), align="justify", style="MT Trong hộp"))
            self.out.append(Box(inner, frame=self.col("brand!35"), frame_pt=0.4, pad_pt=5))
        if self.show and (b.solution or b.answer):
            f = Fmt(size=11, color=self.col("brand!85!black"))
            inner = [P([T("Lời giải.", replace(f, bold=True, sans=True)), T("\u00A0")]
                       + (self.inl(b.answer, where, replace(f, bold=True)) if b.answer else []),
                       align="left")]
            if b.solution:
                inner += self.blocks(b.solution, where, f)
            self.out.append(Box(inner, frame=self.col("brand!55"), left_only_pt=2, pad_pt=5))

    def wrap_problem(self, stmt: str, lead: list, where: str) -> list:
        """Hộp hình treo phải [[wrap]] của PDF (\\hangindent −0,38\\linewidth) ⇒ bảng 2 cột:
        đề trái 0,62 · hình phải 0,36 (y chỗ đặt của PDF)."""
        head, _, tail = stmt.partition("[[/wrap]]")
        box = head.split("[[wrap]]", 1)[-1]
        pre = head.split("[[wrap]]", 1)[0] if "[[wrap]]" in head else ""
        tk = re.search(r"\\begin\{tikzpicture\}.*?\\end\{tikzpicture\}", box, re.S)
        cap = re.search(r"\\color\{muted\}([^{}]*)\}", box)
        mh = re.search(r"max totalheight=([\d.]+cm)", box)
        left = self.blocks(pre + tail, where)
        _prepend(left, lead)
        fig = Fig(snippet=tk.group(0) if tk else "", caption=[T(cap.group(1).strip(), Fmt())] if cap else [],
                  max_w_frac=1.0, max_h_cm=float(mh.group(1)[:-2]) if mh else 3.2)
        if not tk:
            self.ctx.where = where
            self.ctx.lạ("hộp [[wrap]] không có tikzpicture")
            return left
        return [Tbl([[Cell(left), Cell([fig])]], widths=[0.62, 0.36], borders=False, cell_pad_pt=0,
                    align="left")]

    def table(self, b, where):
        if b.caption:
            self.out.append(P(self.inl(b.caption, where, Fmt(bold=True, sans=True, size=10,
                                                               color=self.C["brand"])),
                              align="left", keep_next=True, space_after=2))
        ncol = len(b.headers) or (len(b.rows[0]) if b.rows else 1)
        rows = []
        if b.headers:
            rows.append([Cell([P(self.inl(h, where, Fmt(bold=True)), align="left" if k == 0 else "center")])
                         for k, h in enumerate(b.headers)])
        for r in b.rows:
            rows.append([Cell([P(self.inl(c, where), align="left" if k == 0 else "center")])
                         for k, c in enumerate(r)])
        self.out.append(Tbl(rows, widths=None, borders=True, header_rows=1 if b.headers else 0,
                            cell_pad_pt=3, border_color=self.C["ink"]))

    def mindmap(self, b, where):
        if b.caption:
            self.out.append(P(self.inl(b.caption, where, Fmt(bold=True, sans=True, size=10,
                                                               color=self.C["brand"])),
                              align="left", keep_next=True, space_after=2))

        def node(n) -> str:
            return "[{" + _texify(n.label, self.show) + "}\n" + "".join(node(c) for c in n.children) + "]\n"
        code = ("\\begin{forest}\n[{" + _texify(b.root, self.show) + "}, mindmap\n"
                + "".join(node(br) for br in b.branches) + "]\n\\end{forest}")
        self.out.append(Fig(snippet=code, max_w_frac=1.0, max_h_cm=8.5))

    def figure_block(self, b, where) -> list:
        cap = self.inl(b.caption, where, Fmt(italic=True, sans=True, size=9, color=self.C["muted"])) \
            if b.caption else []
        frac = None
        if b.width:
            m = re.match(r"\s*([\d.]+)\s*\\(?:linewidth|textwidth)", b.width)
            frac = float(m.group(1)) if m else None
        if b.image:
            return [Fig(image=self.img_path(b.image), caption=cap, force_w_frac=frac or 0.55)]
        return [Fig(snippet=b.tikz, caption=cap, max_w_frac=frac or 0.9)]


def _stage_color_name(stage) -> str:
    return {"review": "stage1", "concept": "stage2", "practice1": "stage3",
            "practice2": "stage4", "reflection": "stage5"}[stage.kind]


def _plain_label(label: str) -> str:
    return re.sub(r"\\[a-zA-Z]+|[{}$]", "", label or "").strip()


def _prepend(blocks: list, runs: list):
    """Gắn nhãn 'Bài N. ★★☆' vào đầu đoạn đầu tiên của đề (đề mở bằng bảng/hình ⇒ đoạn riêng)."""
    if blocks and isinstance(blocks[0], P):
        blocks[0].inl[:0] = runs
    elif blocks and isinstance(blocks[0], Tbl) and blocks[0].rows and blocks[0].rows[0] \
            and blocks[0].rows[0][0].blocks and isinstance(blocks[0].rows[0][0].blocks[0], P):
        blocks[0].rows[0][0].blocks[0].inl[:0] = runs           # minipage đề trái – hình phải
    else:
        blocks.insert(0, P(list(runs), align="left", keep_next=True))


def _texify_light(s: str) -> str:
    """Chỉ bỏ thẻ mức độ [NB]… — các token [[…]] để docx_text tự hiểu (giữ được bề rộng ô)."""
    return re.sub(r"\[(?:NB|TH|VD|VDC)\]\s*", "", s or "")


# ═══════════════════════════════ 2. Gom tài nguyên ═══════════════════════════════
def _walk(blocks):
    for b in blocks:
        yield b
        if isinstance(b, Tbl):
            for row in b.rows:
                for c in row:
                    yield from _walk(c.blocks)
        elif isinstance(b, Box):
            yield from _walk(b.blocks)


def _inlines_of(b) -> list:
    if isinstance(b, P):
        return b.inl
    if isinstance(b, Fig):
        return b.caption
    return []


def collect(blocks) -> tuple[list[str], list[str]]:
    maths, snippets = [], []
    for b in _walk(blocks):
        if isinstance(b, DMath):
            maths.append(b.tex)
        if isinstance(b, Fig) and b.snippet:
            snippets.append(b.snippet)
        for x in _inlines_of(b):
            if isinstance(x, M):
                maths.append(x.tex)
            elif isinstance(x, Img) and x.snippet:
                snippets.append(x.snippet)
    return maths, snippets


# ═══════════════════════════════ 3. IR → Word ═══════════════════════════════
class Emitter:
    def __init__(self, doc, maths: dict, figs: dict, report: Report, colors: dict):
        self.doc = doc
        self.maths = maths
        self.figs = figs
        self.r = report
        self.C = colors
        self._qr: dict[str, str] = {}

    # ---------- khối ----------
    def blocks(self, cont, blocks: list, width_in: float, in_box: bool = False):
        for b in blocks:
            self.block(cont, b, width_in, in_box)

    def block(self, cont, b, w: float, in_box: bool):
        if isinstance(b, P):
            self.para(cont, b, w, in_box)
        elif isinstance(b, DMath):
            self.dmath(cont, b)
        elif isinstance(b, Fig):
            self.fig(cont, b, w)
        elif isinstance(b, Tbl):
            self.table(cont, b, w)
        elif isinstance(b, Box):
            self.box(cont, b, w)
        elif isinstance(b, PageBreak):
            cont.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(b, WriteLines):
            self.writelines(cont, b, w)

    def para(self, cont, b: P, w: float, in_box: bool):
        p = cont.add_paragraph(style=b.style or ("MT Trong hộp" if in_box else None))
        pf = p.paragraph_format
        pf.alignment = _ALIGN[b.align]
        if b.tight:
            pf.space_after = Pt(0)
        if b.space_after is not None:
            pf.space_after = Pt(b.space_after)
        if b.space_before is not None:
            pf.space_before = Pt(b.space_before)
        if b.indent_pt:
            pf.left_indent = Pt(b.indent_pt)
        if b.hanging_pt:
            pf.first_line_indent = Pt(-b.hanging_pt)
        if b.keep_next:
            pf.keep_with_next = True
        if b.border_bottom:
            X.ppr_set(p._p, parse_xml(X.border_xml("pBdr", {"bottom": (0.7, b.border_bottom, 1)})))
        tabs = [x for x in b.inl if isinstance(x, Tab)]
        if tabs:
            pos = Inches(max(0.5, w - 0.02))
            pf.tab_stops.add_tab_stop(pos, WD_TAB_ALIGNMENT.RIGHT,
                                      WD_TAB_LEADER.DOTS if tabs[-1].dots else WD_TAB_LEADER.SPACES)
        self.inlines(p, b.inl)
        return p

    def inlines(self, p, inl: list):
        # Công thức cỡ theo chữ của đoạn — CHỈ khi mọi run chữ thật (có chữ/số) đều đặt cỡ
        # (ô bảng TM 9,5pt). Sao ★ cỡ 10 ở nhãn bài không được kéo công thức đề bài nhỏ theo.
        words = [x for x in inl if isinstance(x, T) and any(ch.isalnum() for ch in x.text)]
        self._math_pt = (min(x.fmt.size for x in words)
                         if words and all(x.fmt.size for x in words) else None)
        for x in inl:
            if isinstance(x, T):
                self.run(p, x)
            elif isinstance(x, M):
                self.math(p, x.tex)
            elif isinstance(x, Tab):
                p.add_run().add_tab()
            elif isinstance(x, Img):
                path = self.img(x)
                if path:
                    p.add_run().add_picture(path, height=Pt(x.h_pt))

    def run(self, p, t: T):
        r = p.add_run(t.text)
        f = t.fmt
        if f.bold:
            r.bold = True
        if f.italic:
            r.italic = True
        if f.underline:
            r.underline = True
        if f.smallcaps:
            r.font.small_caps = True
        if f.sup:
            r.font.superscript = True
        if f.color:
            r.font.color.rgb = RGBColor.from_string(f.color)
        if f.size:
            r.font.size = Pt(f.size)
        if f.sans:
            _set_font(r, SANS)
        if f.shade:
            X.rpr_set(r._r, X.shd(f.shade))
        if f.boxed:
            X.rpr_set(r._r, parse_xml(f'<w:bdr {nsdecls("w")} w:val="single" w:sz="4" w:space="0" w:color="auto"/>'))
        return r

    def math(self, p, tex: str):
        if is_only_space(tex):
            p.add_run("\u2003")
            return
        el, why = self.maths.get(tex, (None, "chưa đổi"))
        if el is None:
            self.r.issues.append(f"công thức hỏng ({why}): ${tex}$")
            p.add_run(f"${tex}$")
            return
        el = copy.deepcopy(el)
        if getattr(self, "_math_pt", None):
            sz = round(self._math_pt * 2)
            for mr in el.iter(f"{{{M_NS}}}r"):
                rpr = parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>')
                # trong m:r, w:rPr đứng SAU m:rPr và TRƯỚC m:t
                mt = mr.find(f"{{{M_NS}}}t")
                if mt is not None:
                    mt.addprevious(rpr)
        p._p.append(el)

    def dmath(self, cont, b: DMath):
        p = cont.add_paragraph()
        el, why = self.maths.get(b.tex, (None, "chưa đổi"))
        if el is None:
            self.r.issues.append(f"công thức hỏng ({why}): \\[{b.tex}\\]")
            p.add_run(b.tex)
            return
        para = parse_xml(f'<m:oMathPara xmlns:m="{M_NS}"><m:oMathParaPr><m:jc m:val="{b.align if b.align != "justify" else "center"}"/></m:oMathParaPr></m:oMathPara>')
        para.append(copy.deepcopy(el))
        p._p.append(para)

    # ---------- hình ----------
    def img(self, x: Img) -> str | None:
        if x.path.startswith("qr:"):
            return self.qr(x.path[3:])
        if x.path:
            return x.path
        res = self.figs.get(x.snippet)
        if res is None or isinstance(res, Exception):
            self.r.issues.append(f"hình nhỏ hỏng: {res}")
            return None
        return str(res.png)

    def qr(self, url: str) -> str:
        if url not in self._qr:
            import hashlib
            import segno
            d = settings.ROOT / "storage" / "cache" / "qr"
            d.mkdir(parents=True, exist_ok=True)
            f = d / (hashlib.sha1(url.encode()).hexdigest() + ".png")
            if not f.exists():
                segno.make(url, error="m").save(str(f), scale=10, border=1)
            self._qr[url] = str(f)
        return self._qr[url]

    def fig(self, cont, b: Fig, w: float):
        if b.image:
            path = b.image
            if not Path(path).exists():
                self.r.issues.append(f"thiếu ảnh: {path}")
                return
            from PIL import Image
            with Image.open(path) as im:
                dpi = (im.info.get("dpi") or (150, 150))[0] or 150
                nat_w, nat_h = im.width / dpi, im.height / dpi
        else:
            res = self.figs.get(b.snippet)
            if res is None or isinstance(res, Exception):
                self.r.issues.append(f"hình hỏng: {res} ← {b.snippet[:80]!r}")
                return
            path, nat_w, nat_h = str(res.png), res.w_in, res.h_in
        if b.force_w_frac:
            tw = w * b.force_w_frac
        else:
            tw = min(nat_w, w * b.max_w_frac)
        th = nat_h * tw / nat_w
        if b.max_h_cm and th > b.max_h_cm / 2.54:
            th = b.max_h_cm / 2.54
            tw = nat_w * th / nat_h
        p = cont.add_paragraph(style="MT Hình")
        p.paragraph_format.alignment = _ALIGN[b.align]
        p.add_run().add_picture(path, width=Inches(tw))
        if b.caption:
            cp = cont.add_paragraph(style="MT Chú thích")
            for t in b.caption:
                if isinstance(t, T):
                    t = T(t.text, replace(t.fmt, italic=True, sans=True, size=9, color=self.C["muted"]))
                self.inlines(cp, [t])

    # ---------- bảng / hộp ----------
    def table(self, cont, b: Tbl, w: float):
        ncol = max(sum(c.span for c in r) for r in b.rows) if b.rows else 1
        if b.widths:
            ncol = max(ncol, len(b.widths))
        widths = b.widths
        if widths is None:
            widths = _auto_widths(b, ncol)
        widths = [x * w for x in widths[:ncol]] + [w / ncol] * (ncol - len(widths))
        tot = sum(widths)
        if tot > w:
            widths = [x * w / tot for x in widths]
        t = cont.add_table(rows=len(b.rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER if b.align == "center" else WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        tblPr = t._tbl.tblPr
        for el in tblPr.findall(qn("w:tblStyle")):
            tblPr.remove(el)
        if b.borders:
            X.tblpr_set(t._tbl, parse_xml(X.border_xml("tblBorders", {s: (0.5, b.border_color, 0) for s in
                                                                     ("top", "left", "bottom", "right", "insideH", "insideV")})))
        else:
            X.tblpr_set(t._tbl, parse_xml(X.nil_borders_xml("tblBorders")))
        pad = round(b.cell_pad_pt * 20)
        vpad = 10 if b.compact else pad
        X.tblpr_set(t._tbl, parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="{vpad}" w:type="dxa"/>'
                               f'<w:left w:w="{max(pad, 60)}" w:type="dxa"/><w:bottom w:w="{vpad}" w:type="dxa"/>'
                               f'<w:right w:w="{max(pad, 60)}" w:type="dxa"/></w:tblCellMar>'))
        # Bề rộng tuyệt đối (dxa) cho TỪNG ô — tex2gdoc: bảng chỉ có lưới thì vào Word/Docs bị bóp.
        grid = t._tbl.tblGrid
        for gc, wd in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(round(wd * 1440)))
        taken: set[tuple[int, int]] = set()          # ô lưới đã bị ô gộp dọc phía trên chiếm
        for ri, row in enumerate(b.rows):
            ci = 0
            for cell in row:
                while (ri, ci) in taken:
                    ci += 1
                if ci >= ncol:
                    break
                tc = t.cell(ri, ci)
                last_r = min(len(b.rows) - 1, ri + cell.vspan - 1)
                if cell.span > 1 or last_r > ri:
                    tc = tc.merge(t.cell(last_r, min(ncol - 1, ci + cell.span - 1)))
                    for rr in range(ri + 1, last_r + 1):
                        for cc in range(ci, ci + cell.span):
                            taken.add((rr, cc))
                cw = sum(widths[ci:ci + cell.span])
                tc.width = Inches(cw)
                if cell.shade:
                    X.tcpr_set(tc._tc, X.shd(cell.shade))
                for extra in tc.paragraphs[1:]:            # merge() để lại đoạn rỗng của ô bị gộp
                    if not extra.text and not extra._p.findall(qn("w:r")):
                        extra._p.getparent().remove(extra._p)
                first = tc.paragraphs[0]
                inner_w = max(0.3, cw - 2 * max(pad, 60) / 1440)
                self.blocks(tc, cell.blocks, inner_w, in_box=True)
                if b.compact:
                    for pp in tc.paragraphs:
                        pp.paragraph_format.space_after = Pt(0)
                        pp.paragraph_format.space_before = Pt(0)
                        pp.paragraph_format.line_spacing = 0.95   # kiểu Multiple: dòng có phân số vẫn tự giãn
                if len(tc.paragraphs) > 1 or tc.tables:
                    first._p.getparent().remove(first._p)
                if ri < b.header_rows:
                    for pp in tc.paragraphs:
                        for rr in pp.runs:
                            rr.bold = True
                ci += cell.span
            if ri < b.header_rows:
                trPr = t.rows[ri]._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        _after_table(cont)

    def box(self, cont, b: Box, w: float):
        t = cont.add_table(rows=1, cols=1)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        for el in tblPr.findall(qn("w:tblStyle")):
            tblPr.remove(el)
        if b.left_only_pt:
            X.tblpr_set(t._tbl, parse_xml(X.border_xml("tblBorders", {"left": (b.left_only_pt, b.frame, 0)})))
        else:
            X.tblpr_set(t._tbl, parse_xml(X.border_xml("tblBorders", {s: (b.frame_pt, b.frame, 0) for s in
                                                                     ("top", "left", "bottom", "right")})))
        pad = round(b.pad_pt * 20)
        X.tblpr_set(t._tbl, parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:top w:w="{pad}" w:type="dxa"/>'
                               f'<w:left w:w="{pad + 40}" w:type="dxa"/><w:bottom w:w="{pad}" w:type="dxa"/>'
                               f'<w:right w:w="{pad + 40}" w:type="dxa"/></w:tblCellMar>'))
        t._tbl.tblGrid.findall(qn("w:gridCol"))[0].set(qn("w:w"), str(round(w * 1440)))
        tc = t.cell(0, 0)
        tc.width = Inches(w)
        if b.back:
            X.tcpr_set(tc._tc, X.shd(b.back))
        first = tc.paragraphs[0]
        if b.title:
            first.style = self.doc.styles["MT Nhãn hộp"]
            self.run(first, T(f"\u00A0{b.title}\u00A0", Fmt(bold=True, sans=True, size=10,
                                                           color=b.title_fg, shade=b.title_bg)))
        inner_w = w - 2 * (pad + 40) / 1440
        self.blocks(tc, b.blocks, inner_w, in_box=True)
        if not b.title and (len(tc.paragraphs) > 1 or tc.tables):
            first._p.getparent().remove(first._p)
        _after_table(cont)

    def writelines(self, cont, b: WriteLines, w: float):
        if b.count == 0:
            p = cont.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            return
        if b.variant == "draw":
            t = cont.add_table(rows=1, cols=1)
            X.tblpr_set(t._tbl, parse_xml(f'<w:tblBorders {nsdecls("w")}>' + "".join(
                f'<w:{s} w:val="dashed" w:sz="4" w:space="0" w:color="{self.C["rule"]}"/>'
                for s in ("top", "left", "bottom", "right")) + "</w:tblBorders>"))
            row = t.rows[0]
            row.height = Inches(b.count / 2.54)
            p = t.cell(0, 0).paragraphs[0]
            self.run(p, T("Vẽ hình vào khung này", Fmt(italic=True, sans=True, size=8, color=self.C["muted"])))
            _after_table(cont)
            return
        for _ in range(b.count):
            p = cont.add_paragraph()
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(w - 0.05), WD_TAB_ALIGNMENT.RIGHT,
                                                      WD_TAB_LEADER.DOTS)
            p.add_run().add_tab()


def _set_font(run, font: str):
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = parse_xml(f"<w:rFonts {nsdecls('w')}/>")
        rpr.insert(0, rf)
    for k in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{k}"), font)


def _after_table(cont):
    """Hai bảng liền nhau Word sẽ GỘP làm một ⇒ chèn một đoạn mỏng ngăn cách."""
    p = cont.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)
    pf.line_spacing = Pt(4)
    X.ppr_set(p._p, parse_xml(f'<w:rPr {nsdecls("w")}><w:sz w:val="6"/></w:rPr>'))


def _plain_len(blocks) -> int:
    n = 0
    for b in _walk(blocks):
        for x in _inlines_of(b):
            if isinstance(x, T):
                n += len(x.text)
            elif isinstance(x, M):
                n += max(1, int(len(re.sub(r"\\[a-zA-Z]+|[{}^_]", "", x.tex)) * 0.8))
    return n


def _auto_widths(b: Tbl, ncol: int) -> list[float]:
    """Bảng `l|c|c…` không khai bề rộng: chia theo độ dài chữ dài nhất mỗi cột (như LaTeX tự
    co theo nội dung), kẹp trong [0,08; 0,5] bề ngang, tổng ≤ 1."""
    mx = [1] * ncol
    for row in b.rows:
        ci = 0
        for c in row:
            if c.span == 1 and ci < ncol:
                mx[ci] = max(mx[ci], _plain_len(c.blocks))
            ci += c.span
    est = [min(0.5, max(0.08, (m + 2) * 0.0115)) for m in mx]
    s = sum(est)
    return [e / s for e in est] if s > 1 else est


# ═══════════════════════════════ API ═══════════════════════════════
def ensure_template(path: Path, ngang: bool = False) -> Path:
    """Máy nào thiếu khuôn (clone thiếu, lỡ xoá) thì tự dựng lại từ design_tokens."""
    if not path.exists():
        from scripts.make_docx_template import build
        build(path, ngang=ngang)
    return path



def render_lesson_docx(lesson, lesson_dir: Path | None = None, show_solution: bool = False):
    """→ (Document, Report). Không ghi file — `build_docx` lo ghi sau khi cổng duyệt."""
    from src.compiler.tikz_raster import render_many
    tokens = load_tokens()
    set_palette(tokens["colors"])
    ir = LessonIR(lesson, tokens, show_solution, lesson_dir)
    blocks = ir.build()
    report = Report()
    report.issues += [f"lệnh chưa có luật đổi sang Word — {u}" for u in ir.ctx.unknown]
    report.counts["problems"] = ir.n_problem
    maths, snippets = collect(blocks)
    report.counts["maths"] = sum(1 for m in maths if not is_only_space(m))
    mres = convert_many([m for m in maths if not is_only_space(m)])
    fres = render_many(snippets) if snippets else {}
    doc = Document(str(ensure_template(TEMPLATE)))
    em = Emitter(doc, mres, fres, report, tokens["colors"])
    em.blocks(doc, blocks, TEXT_W_IN)
    _core_props(doc, lesson)
    return doc, report


def _core_props(doc, lesson):
    cp = doc.core_properties
    cp.title = lesson.title
    cp.subject = lesson.grade_label or ""
    cp.author = "MathTech"
    cp.comments = "Sinh tự động từ seed JSON — sửa nội dung ở JSON để bản PDF/Word luôn khớp."
